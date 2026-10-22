"""Day 9 作业一：按正文原顺序读取 Word 段落与简单表格。"""

from pathlib import Path
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml.etree import XMLSyntaxError

from app.pdf_ingestion import DocumentParseError


def read_docx_blocks(path: str | Path) -> list[dict]:
    """本轮只实现此函数；打开文件及受限结构校验由 open_docx_document 提供。

    输入：真实 .docx 文件路径。输出：正文中的段落/表格列表，保持交错顺序。
    1. document = open_docx_document(path)。
    2. enumerate(document.iter_inner_content(), start=1) 遍历正文元素。
       block 是原正文元素序号，包含后来跳过的空段落，不按输出列表长度重编号。
    3. isinstance(element, Paragraph) 时：
       text = element.text.strip()；空文字段落跳过，否则保存
       {'kind': 'paragraph', 'block': 原序号, 'style': element.style.name, 'text': text}。
       保留段落内部换行；Title/Heading 1 等样式名称原样保留，供后续判断章节。
    4. isinstance(element, Table) 时：
       按 element.rows -> row.cells 读取，保留二维结构及表头：
       {'kind': 'table', 'block': 原序号, 'rows': [[cell.text.strip(), ...], ...]}。
       空单元格保留 ''，不能丢掉后让列错位，也不提前压平成一串文字。
    5. 整份没有有效块时抛 DocumentParseError，不返回 [] 当作成功空文档。

    不能分别循环 document.paragraphs 和 document.tables 后拼接，这会破坏原顺序。
    此轮不猜页码，不切分/编码/入库，也不创建新 Word 文件或改写输入。
    支持普通正文段落、样式标题和无合并的矩形表；复杂结构 helper 会明确拒绝。
    下一课再把表头/数据行转为可检索事实，并记录所在章节和正文块位置。
    """
    document = open_docx_document(path)
    results = []
    for block, element in enumerate(document.iter_inner_content(),start=1):
        if isinstance(element, Paragraph):
            text = element.text.strip()
            if not text:
                continue
            results.append({
                "kind": "paragraph",
                "block": block,
                "style": element.style.name,
                "text": text
            })
        if isinstance(element, Table):
            rows = [
                [cell.text.strip() for cell in row.cells]
                for row in element.rows
            ]
            results.append({
                "kind": "table",
                "block": block,
                "rows": rows
            })
    if not results:
        raise DocumentParseError("整份文档为空")
    return results


def open_docx_document(path: str | Path):
    """助教提供：本课受限 DOCX 适配，失败不伪装成空文档或完整快照。"""
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(path)
    if path.suffix.lower() != ".docx":
        raise DocumentParseError("本课仅接收 .docx，不处理旧版 .doc")
    try:
        document = Document(str(path))
    except (PackageNotFoundError, BadZipFile, XMLSyntaxError, KeyError, ValueError) as error:
        raise DocumentParseError("Word 文件损坏或不是可解析的 DOCX") from error
    # 避免只提取部分正文却宣称文档已完整解析。此处不是完整 Word 特性识别器。
    unsupported = {qn(f"w:{name}") for name in ("ins", "del", "sdt", "txbxContent", "altChunk", "drawing", "pict")}
    if any(element.tag in unsupported for element in document.element.body.iter()):
        raise DocumentParseError("本课不处理修订、内容控件、文本框、图片或外部正文，请先人工整理")
    for table in document.tables:
        if not table.rows or not table.columns:
            raise DocumentParseError("表格缺少行或列")
        width = len(table.columns)
        for row in table.rows:
            if row.grid_cols_before or row.grid_cols_after or len(row.cells) != width:
                raise DocumentParseError("本课只处理完整矩形表格，不处理缺失网格列")
            for cell in row.cells:
                properties = cell._tc.tcPr
                if cell.grid_span != 1 or (properties is not None and properties.find(qn("w:vMerge")) is not None):
                    raise DocumentParseError("本课不处理合并单元格，不能按重复单元格猜测含义")
                if cell.tables:
                    raise DocumentParseError("本课不处理嵌套表格")
        if not any(cell.text.strip() for row in table.rows for cell in row.cells):
            raise DocumentParseError("表格完全为空，停止导入")
    return document
