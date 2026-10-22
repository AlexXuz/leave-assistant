"""助教样例生成器：使用 Codex 附带 Python，生成正文/表格交错的教学 DOCX。"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUTPUT = Path(__file__).resolve().parents[1] / "data/ingestion_lab/travel-policy.docx"


def set_style_font(style, size):
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")


def add_table(document, headers, records, widths):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)
    for index, text in enumerate(headers):
        table.rows[0].cells[index].text = text
    for record in records:
        for cell, text in zip(table.add_row().cells, record):
            cell.text = text
    for row_number, row in enumerate(table.rows):
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side, value in (("top", 100), ("bottom", 100), ("left", 130), ("right", 130)):
                element = OxmlElement(f"w:{side}")
                element.set(qn("w:w"), str(value))
                element.set(qn("w:type"), "dxa")
                margins.append(element)
            properties.append(margins)
            borders = OxmlElement("w:tcBorders")
            for side in ("top", "bottom", "left", "right"):
                border = OxmlElement(f"w:{side}")
                for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    border.set(qn(f"w:{key}"), value)
                borders.append(border)
            properties.append(borders)
            if row_number == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "E8EEF5")
                properties.append(shading)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.2
                for run in paragraph.runs:
                    run.font.size = Pt(10.5)
                    run.bold = row_number == 0
    header_flag = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(header_flag)
    return table


def main():
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(.85)
    section.top_margin = section.bottom_margin = Inches(.75)
    for name, size in (("Normal", 11), ("Title", 21), ("Heading 1", 14)):
        set_style_font(document.styles[name], size)
    # 清除默认模板继承的标题装饰边框，保持本课样例的纯文字标题。
    for style in document.styles:
        for border in style.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    normal = document.styles["Normal"].paragraph_format
    normal.line_spacing = 1.35
    normal.space_after = Pt(9)
    document.styles["Title"].paragraph_format.space_after = Pt(10)
    heading = document.styles["Heading 1"].paragraph_format
    heading.space_before = Pt(16)
    heading.space_after = Pt(8)
    document.add_paragraph("青禾实验室差旅制度", "Title")
    metadata = document.add_paragraph("文档 ID：travel-policy　版本：2026-10-22-v1")
    for run in metadata.runs:
        run.font.size = Pt(9)
    document.add_paragraph("本制度说明差旅住宿标准、例外审批及报销凭证要求，供员工办理出差申请和住宿报销时查阅。")
    document.add_heading("住宿标准", level=1)
    document.add_paragraph("以下金额为每人每晚住宿报销上限。")
    add_table(document, ("城市", "住宿报销上限", "计量单位"),
              (("北京和上海", "520 元", "每人每晚"), ("其他城市", "380 元", "每人每晚")), (2.6, 2.1, 2.1))
    after_table = document.add_paragraph("超过上限的住宿，需要在预订前取得部门经理的书面批准。")
    after_table.paragraph_format.space_before = Pt(9)
    document.add_heading("申请与凭证", level=1)
    document.add_paragraph("员工应在出差前提交差旅申请，并由直属主管审批。")
    add_table(document, ("办理事项", "所需材料"), (("住宿报销", "对应的住宿发票和订单记录"),), (2.0, 4.8))
    closing = document.add_paragraph("本文件未规定每日餐费补贴标准，也未提供员工个人报销记录。")
    closing.paragraph_format.space_before = Pt(9)
    document.core_properties.title = "青禾实验室差旅制度"
    document.core_properties.author = "青禾实验室"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
