"""助教验收：真实 DOCX 的交错顺序、表格结构和解析失败；无模型/数据库调用。"""

import argparse
import hashlib
import json
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from app import docx_ingestion as exercise
from app.pdf_ingestion import DocumentParseError

SAMPLE = Path(__file__).resolve().parents[1] / "data/ingestion_lab/travel-policy.docx"


def check():
    before = hashlib.sha256(SAMPLE.read_bytes()).hexdigest()
    blocks = exercise.read_docx_blocks(SAMPLE)
    assert hashlib.sha256(SAMPLE.read_bytes()).hexdigest() == before, "不能改写原文件"
    assert len(blocks) == 11, "样例应为11个正文块，表格不应被丢弃或拆成段落"
    assert [block["block"] for block in blocks] == list(range(1, 12)), "保留正文原始元素序号"
    assert [block["kind"] for block in blocks] == ["paragraph"] * 5 + ["table", "paragraph", "paragraph", "paragraph", "table", "paragraph"], \
        "段落和表格必须保持原来的交错顺序"
    for block in blocks:
        if block["kind"] == "paragraph":
            assert set(block) == {"kind", "block", "style", "text"}
        else:
            assert set(block) == {"kind", "block", "rows"}
    assert blocks[0]["style"] == "Title"
    assert blocks[3]["style"] == blocks[7]["style"] == "Heading 1"
    assert blocks[5]["rows"] == [["城市", "住宿报销上限", "计量单位"], ["北京和上海", "520 元", "每人每晚"], ["其他城市", "380 元", "每人每晚"]]
    assert blocks[6]["text"] == "超过上限的住宿，需要在预订前取得部门经理的书面批准。"
    assert blocks[9]["rows"] == [["办理事项", "所需材料"], ["住宿报销", "对应的住宿发票和订单记录"]]
    print("PASS：真实中文 DOCX，段落/表格交错顺序、样式与表头/数值/单位保持正确。")
    with TemporaryDirectory(prefix="day9-docx-") as folder:
        directory = Path(folder)
        document = Document()
        document.add_paragraph(" \n ")
        document.add_heading(" 第一节 ", 1)
        table = document.add_table(rows=2, cols=2)
        for cell, value in zip(table.rows[0].cells, ("项目", "备注")):
            cell.text = value
        table.rows[1].cells[0].text = "甲\n乙"
        table.rows[1].cells[1].text = ""
        document.add_paragraph(" 表后条件 ")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "第二个表"
        document.add_paragraph("结束")
        path = directory / "mixed.docx"
        document.save(path)
        result = exercise.read_docx_blocks(path)
        assert result == [
            {"kind": "paragraph", "block": 2, "style": "Heading 1", "text": "第一节"},
            {"kind": "table", "block": 3, "rows": [["项目", "备注"], ["甲\n乙", ""]]},
            {"kind": "paragraph", "block": 4, "style": "Normal", "text": "表后条件"},
            {"kind": "table", "block": 5, "rows": [["第二个表"]]},
            {"kind": "paragraph", "block": 6, "style": "Normal", "text": "结束"},
        ], "跳过空段落后仍保留原块序号；保留空单元格和内部换行"
        for mode in ("empty", "merged", "nested", "empty-table", "broken"):
            document = Document()
            document.add_paragraph(" \n ")
            if mode == "merged":
                table = document.add_table(rows=2, cols=2)
                table.cell(0, 0).merge(table.cell(0, 1)).text = "合并表头"
            elif mode == "nested":
                document.add_table(rows=1, cols=1).cell(0, 0).add_table(rows=1, cols=1)
            elif mode == "empty-table":
                document.add_table(rows=1, cols=1)
            path = directory / f"{mode}.docx"
            if mode == "broken":
                path.write_bytes(b"not a docx")
            else:
                document.save(path)
            try:
                exercise.read_docx_blocks(path)
            except DocumentParseError:
                pass
            else:
                raise AssertionError(f"{mode} 不能当作完整有效文档成功返回")
        try:
            exercise.read_docx_blocks(directory / "missing.docx")
        except FileNotFoundError:
            pass
        else:
            raise AssertionError("文件缺失应传出，不转换为空快照")
    print("PASS：不同结构、空段落/单元格、换行以及空白/损坏/不支持结构的失败契约。")
    return blocks


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--show", action="store_true")
    args = parser.parse_args()
    try:
        blocks = check()
    except NotImplementedError as error:
        print(f"待完成：{error}")
        return 2
    except (AssertionError, TypeError, KeyError) as error:
        print(f"FAIL：{error}")
        return 1
    if args.show:
        print(json.dumps(blocks, ensure_ascii=False, indent=2))
    print("本轮仅解析结构；尚未切分/入库，未调用模型；没有伪造 Word 页码。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
